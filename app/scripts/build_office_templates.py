"""Reproduce the empty, synthetic OOXML authoring templates; not a runtime dependency.

Run with python-docx 1.2.0, openpyxl 3.1.5 and python-pptx 1.0.2. Templates
contain no company data, credentials, embedded fonts, macros or executable code.
"""
from datetime import datetime
from io import BytesIO
import hashlib
import importlib.metadata
import json
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / 'app/backend/internal/files/office_templates'
VERSIONS = {'python-docx': '1.2.0', 'openpyxl': '3.1.5', 'python-pptx': '1.0.2'}


def canonical_package(content):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, 'w') as target:
        for name in sorted(source.namelist()):
            if 'thumbnail' in name.lower():
                continue
            data = source.read(name)
            if name.endswith('.rels'):
                tree = ET.fromstring(data)
                for item in list(tree):
                    if item.attrib.get('Type', '').endswith('/metadata/thumbnail'):
                        tree.remove(item)
                data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            if name == '[Content_Types].xml':
                tree = ET.fromstring(data)
                for item in list(tree):
                    if 'thumbnail' in item.attrib.get('PartName', '').lower():
                        tree.remove(item)
                data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            if name == 'docProps/core.xml':
                tree = ET.fromstring(data)
                for item in tree:
                    if item.tag.rsplit('}', 1)[-1] in ('created', 'modified'):
                        item.text = '2000-01-01T00:00:00Z'
                data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, data)
    return output.getvalue()


def main():
    actual = {name: importlib.metadata.version(name) for name in VERSIONS}
    if actual != VERSIONS:
        raise SystemExit('Use the explicitly declared template-generator versions.')
    import docx
    import openpyxl
    import pptx
    TARGET.mkdir(parents=True, exist_ok=True)
    documents = {'docx': docx.Document(), 'xlsx': openpyxl.Workbook(), 'pptx': pptx.Presentation()}
    rows = []
    for kind, document in documents.items():
        props = document.properties if kind == 'xlsx' else document.core_properties
        if kind == 'xlsx':
            props.creator, props.lastModifiedBy = 'CopyPaste', 'CopyPaste'
        else:
            props.author, props.last_modified_by = 'CopyPaste', 'CopyPaste'
        props.created = datetime(2000, 1, 1)
        props.modified = datetime(2000, 1, 1)
        props.title = 'CopyPaste empty authoring template'
        buffer = BytesIO()
        document.save(buffer)
        content = canonical_package(buffer.getvalue())
        file = TARGET / ('empty.' + kind)
        file.write_bytes(content)
        rows.append({'file': file.relative_to(ROOT).as_posix(), 'sha256': hashlib.sha256(content).hexdigest(), 'bytes': len(content)})
    record = {'generator_versions': actual, 'templates': rows,
              'scope': 'Empty synthetic templates and inherited OOXML style/layout parts only; all generators are MIT licensed. No fonts, Office binaries or business documents are distributed.'}
    (TARGET / 'manifest.json').write_bytes((json.dumps(record, indent=2) + '\n').encode())
    print(json.dumps(record))


if __name__ == '__main__':
    main()
