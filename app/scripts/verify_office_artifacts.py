"""Independent native OOXML consumers, not Office rendering acceptance."""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys

APP=Path(__file__).resolve().parents[1]

def inspect(folder):
    import docx
    import openpyxl
    import pptx
    folder=Path(folder)
    word=docx.Document(folder/'created.docx')
    assert word.paragraphs[0].text=='交付说明' and word.tables[0].cell(1,1).text=='1200'
    assert docx.Document(folder/'edited.docx').paragraphs[0].text=='新版内容 香港🙂'
    book=openpyxl.load_workbook(folder/'created.xlsx',data_only=False)
    assert book.sheetnames==['开支','备注'] and book['开支']['B1'].value==1200 and book['开支']['B2'].value=='=SUM(B1:B1)'
    edited=openpyxl.load_workbook(folder/'edited.xlsx',data_only=False)
    assert edited['开支']['A1'].value=='新版内容 香港🙂' and edited['开支']['B2'].value=='=SUM(B1:B1)'
    assert openpyxl.load_workbook(folder/'edited.xlsx',data_only=True)['开支']['B2'].value is None
    deck=pptx.Presentation(folder/'created.pptx')
    assert len(deck.slides)==2 and deck.slides[0].shapes[0].text=='项目概况'
    assert pptx.Presentation(folder/'edited.pptx').slides[0].shapes[0].text=='新版内容 香港🙂'
    print(json.dumps({'passed':True,'independent_consumers':{'python-docx':docx.__version__,'openpyxl':openpyxl.__version__,'python-pptx':pptx.__version__},'artifacts':[{'name':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size}for p in sorted(folder.glob('*'))],'rendering_validated':False}))

def main():
    folder=APP/'.local/office-authoring/consumer-fixtures';folder.mkdir(parents=True,exist_ok=True)
    env={**os.environ,'CP_OFFICE_AUTHOR_FIXTURES':str(folder),'GOPROXY':'https://goproxy.cn,https://proxy.golang.org,direct'}
    result=subprocess.run([str(APP/'.tools/go/bin/go.exe'),'test','./internal/files','-run','^TestOfficeAuthor','-count=1','-v'],cwd=APP/'backend',env=env)
    if result.returncode:return result.returncode
    executable=os.environ.get('CP_OFFICE_PYTHON')
    if not executable or not Path(executable).is_file():raise SystemExit('The explicitly pinned native Python with the three independent consumers is required.')
    return subprocess.call([executable,str(Path(__file__).resolve()),'--inspect',str(folder)])

if __name__=='__main__':
    if len(sys.argv)==3 and sys.argv[1]=='--inspect':inspect(sys.argv[2])
    else:raise SystemExit(main())
